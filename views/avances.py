import streamlit as st
import pandas as pd
import io
import re
import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from utils.database import get_active_of, get_todas_piezas, get_avances_nido, save_avances_mixto, get_total_rechazos, get_connection, get_movimientos_area, get_all_ofs, get_personal_prenomina, get_operadores_por_area, get_local_now
from st_aggrid import AgGrid, GridOptionsBuilder, GridUpdateMode, DataReturnMode, JsCode

# Constantes de diseño
RED = "#EC2024"
BLACK = "#111111"
GRAY = "#D2D3D5"
WHITE = "#FFFFFF"

PROCESSES = ["Ingenieria", "Corte", "Rebabeo", "Doblez", "Barrenado", "Pintura", "Liberado", "Empaque"]


def generate_wip_table_excel(df: pd.DataFrame, of_number: str, area_seleccionada: str, total_wip: int = 0) -> bytes:
    """Genera archivo Excel (.xlsx) con formato corporativo SIGRAMA para la tabla WIP del área."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"WIP {area_seleccionada[:25]}"
    
    # Fuentes y estilos
    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    hdr_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    data_bold = Font(name="Calibri", size=10, bold=True)
    tot_font = Font(name="Calibri", size=10, bold=True, color="111111")
    
    # Fills
    title_fill = PatternFill("solid", fgColor="111111")     # Negro corporativo
    red_banner = PatternFill("solid", fgColor="EC2024")     # Rojo corporativo SIGRAMA
    hdr_fill = PatternFill("solid", fgColor="222222")       # Gris oscuro cabecera
    tot_fill = PatternFill("solid", fgColor="E0E0E0")       # Gris totalizador
    
    fill_wip = PatternFill("solid", fgColor="D4EDDA")       # Verde suave (WIP Disp)
    fill_buenas = PatternFill("solid", fgColor="D1ECF1")    # Azul suave (Proc Buenas)
    fill_malas = PatternFill("solid", fgColor="F8D7DA")     # Rojo suave (Proc Malas)
    fill_tot_req = PatternFill("solid", fgColor="E9ECEF")   # Gris claro (Totales OF)
    alt_fill = PatternFill("solid", fgColor="F9F9F9")       # Alternado suave
    
    # Colores de fuente para columnas específicas
    font_proc_ant = Font(name="Calibri", size=10, bold=True, color="004085") # Azul oscuro
    font_rech_ant = Font(name="Calibri", size=10, bold=True, color="721C24") # Rojo vino
    font_pend_tot = Font(name="Calibri", size=10, bold=True, color="C82333") # Rojo alerta
    font_wip_disp = Font(name="Calibri", size=10, bold=True, color="155724") # Verde oscuro
    
    # Alineaciones
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    
    # Bordes
    thin_border = Border(
        left=Side(style="thin", color="D2D3D5"),
        right=Side(style="thin", color="D2D3D5"),
        top=Side(style="thin", color="D2D3D5"),
        bottom=Side(style="thin", color="D2D3D5")
    )
    thick_bottom = Border(
        left=Side(style="thin", color="D2D3D5"),
        right=Side(style="thin", color="D2D3D5"),
        top=Side(style="thin", color="D2D3D5"),
        bottom=Side(style="medium", color="111111")
    )
    
    # 1. Título y Banner superior
    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = f"SIGRAMA — CONTROL DE PRODUCCION | WIP {area_seleccionada.upper()}"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = center_align
    ws.row_dimensions[1].height = 28
    
    ws.merge_cells("A2:K2")
    sub_cell = ws["A2"]
    now_str = get_local_now().strftime("%d/%m/%Y %H:%M")
    sub_cell.value = f"OF: {of_number}   |   Fecha de Reporte: {now_str}   |   Total Piezas WIP Disponibles: {int(total_wip):,}"
    sub_cell.font = subtitle_font
    sub_cell.fill = red_banner
    sub_cell.alignment = center_align
    ws.row_dimensions[2].height = 20
    
    # Fila vacía de separación
    ws.row_dimensions[3].height = 8
    
    # 2. Encabezados de Columnas (Fila 4)
    cols_def = [
        ("of_number", "OF", 18, center_align),
        ("no_pieza", "No. Pieza", 20, center_align),
        ("nombre_pieza", "Descripción", 38, left_align),
        ("total_requeridas", "Totales (OF)", 14, center_align),
        ("Pendiente Disponible", "WIP Disp.", 14, center_align),
        ("terminadas_ant", "Proc. Ant.", 13, center_align),
        ("rechazadas_ant", "Rech. Ant.", 13, center_align),
        ("Terminadas", "Proc. Buenas", 14, center_align),
        ("Rechazos", "Proc. Malas", 13, center_align),
        ("Motivo", "Motivo de rechazo", 24, left_align),
        ("Pendiente Total OF", "Pend. Total", 14, center_align),
    ]
    
    header_row = 4
    ws.row_dimensions[header_row].height = 25
    for col_idx, (_, col_name, col_width, _) in enumerate(cols_def, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    # 3. Filas de Datos
    curr_row = 5
    for _, r in df.iterrows():
        ws.row_dimensions[curr_row].height = 20
        row_fill = alt_fill if curr_row % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
        
        for col_idx, (field_key, _, _, align_style) in enumerate(cols_def, 1):
            val = r.get(field_key, "")
            
            # Cast numéricos a int para que las fórmulas de suma funcionen
            if field_key in ["total_requeridas", "Pendiente Disponible", "terminadas_ant", "rechazadas_ant", "Terminadas", "Rechazos", "Pendiente Total OF"]:
                try:
                    val = int(val) if pd.notna(val) and val != "" else 0
                except:
                    val = 0
            elif pd.isna(val):
                val = ""
                
            cell = ws.cell(row=curr_row, column=col_idx, value=val)
            cell.alignment = align_style
            cell.border = thin_border
            cell.fill = row_fill
            cell.font = data_font
            
            # Estilos condicionales por columna según AgGrid
            if field_key == "total_requeridas":
                cell.fill = fill_tot_req
                cell.font = data_bold
            elif field_key == "Pendiente Disponible":
                cell.fill = fill_wip
                cell.font = font_wip_disp
            elif field_key == "terminadas_ant":
                cell.font = font_proc_ant
            elif field_key == "rechazadas_ant":
                cell.font = font_rech_ant
            elif field_key == "Terminadas":
                cell.fill = fill_buenas
                cell.font = data_bold
            elif field_key == "Rechazos":
                cell.fill = fill_malas
                cell.font = data_bold
            elif field_key == "Pendiente Total OF":
                cell.font = font_pend_tot
                
        curr_row += 1

    # 4. Fila de Totales
    ws.row_dimensions[curr_row].height = 22
    ws.cell(row=curr_row, column=1, value="TOTALES").font = tot_font
    ws.cell(row=curr_row, column=1).alignment = center_align
    ws.cell(row=curr_row, column=1).fill = tot_fill
    ws.cell(row=curr_row, column=1).border = thick_bottom

    for col_idx in range(2, len(cols_def) + 1):
        cell = ws.cell(row=curr_row, column=col_idx)
        cell.fill = tot_fill
        cell.border = thick_bottom
        field_key = cols_def[col_idx - 1][0]
        
        # Calcular sumas si son numéricas
        if field_key in ["total_requeridas", "Pendiente Disponible", "terminadas_ant", "rechazadas_ant", "Terminadas", "Rechazos", "Pendiente Total OF"]:
            col_letter = get_column_letter(col_idx)
            start_row = 5
            end_row = curr_row - 1
            if end_row >= start_row:
                cell.value = f"=SUM({col_letter}{start_row}:{col_letter}{end_row})"
            else:
                cell.value = 0
            cell.font = tot_font
            cell.alignment = center_align
            
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_corte_detalle_excel(of_number: str) -> bytes:
    """Genera archivo Excel (.xlsx) con el Detalle de Piezas pendientes y terminadas en Corte."""
    conn = get_connection()
    if of_number == "Todas":
        query_p = """
            SELECT p.of_number, p.nido, p.no_pieza, p.nombre_pieza, p.cantidad, n.hojas, n.calibre
            FROM piezas p
            JOIN nidos n ON p.of_number=n.of_number AND p.nido=n.nido
            ORDER BY p.of_number, p.nido, p.no_pieza
        """
        df_p = pd.read_sql_query(query_p, conn)
        query_av = """
            SELECT of_number, nido, COUNT(DISTINCT hoja) as hojas_cortadas
            FROM avances
            WHERE area='Corte' AND hoja IS NOT NULL
            GROUP BY of_number, nido
        """
        df_av = pd.read_sql_query(query_av, conn)
        merge_keys = ['of_number', 'nido']
    else:
        query_p = """
            SELECT p.of_number, p.nido, p.no_pieza, p.nombre_pieza, p.cantidad, n.hojas, n.calibre
            FROM piezas p
            JOIN nidos n ON p.of_number=n.of_number AND p.nido=n.nido
            WHERE p.of_number=?
            ORDER BY p.nido, p.no_pieza
        """
        df_p = pd.read_sql_query(query_p, conn, params=(of_number,))
        query_av = """
            SELECT nido, COUNT(DISTINCT hoja) as hojas_cortadas
            FROM avances
            WHERE of_number=? AND area='Corte' AND hoja IS NOT NULL
            GROUP BY nido
        """
        df_av = pd.read_sql_query(query_av, conn, params=(of_number,))
        merge_keys = ['nido']
    conn.close()

    if df_p.empty:
        df_m = pd.DataFrame()
    else:
        df_m = df_p.merge(df_av, on=merge_keys, how='left')
        df_m['hojas_cortadas'] = df_m['hojas_cortadas'].fillna(0).astype(int)
        df_m['hojas'] = df_m['hojas'].fillna(1).astype(int)
        df_m['pzas_planeadas'] = df_m['cantidad'] * df_m['hojas']
        df_m['pzas_cortadas'] = df_m['cantidad'] * df_m[['hojas', 'hojas_cortadas']].min(axis=1)
        df_m['pzas_pendientes'] = df_m['pzas_planeadas'] - df_m['pzas_cortadas']
        
        def calc_estado(r):
            if r['hojas_cortadas'] >= r['hojas']:
                return "Terminado"
            elif r['hojas_cortadas'] > 0:
                return "En Proceso"
            else:
                return "Pendiente"
        df_m['estado'] = df_m.apply(calc_estado, axis=1)

    wb = openpyxl.Workbook()
    
    # Fuentes y estilos
    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    sub_font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    hdr_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    tot_font = Font(name="Calibri", size=10, bold=True, color="111111")
    
    title_fill = PatternFill("solid", fgColor="111111")
    red_banner = PatternFill("solid", fgColor="EC2024")
    hdr_fill = PatternFill("solid", fgColor="222222")
    tot_fill = PatternFill("solid", fgColor="E0E0E0")
    
    fill_term = PatternFill("solid", fgColor="D1ECF1")   # Azul suave
    fill_pend = PatternFill("solid", fgColor="FFF3CD")   # Amarillo suave
    fill_ok   = PatternFill("solid", fgColor="D4EDDA")   # Verde suave
    alt_fill  = PatternFill("solid", fgColor="F9F9F9")
    
    font_term = Font(name="Calibri", size=10, bold=True, color="0C5460")
    font_pend = Font(name="Calibri", size=10, bold=True, color="856404")
    
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    
    thin_border = Border(
        left=Side(style="thin", color="D2D3D5"),
        right=Side(style="thin", color="D2D3D5"),
        top=Side(style="thin", color="D2D3D5"),
        bottom=Side(style="thin", color="D2D3D5")
    )
    thick_bottom = Border(
        left=Side(style="thin", color="D2D3D5"),
        right=Side(style="thin", color="D2D3D5"),
        top=Side(style="thin", color="D2D3D5"),
        bottom=Side(style="medium", color="111111")
    )

    # HOJA 1: DETALLE POR NIDO Y PIEZA
    ws1 = wb.active
    ws1.title = "Detalle Nidos y Piezas"
    
    ws1.merge_cells("A1:K1")
    c1 = ws1["A1"]
    c1.value = "SIGRAMA — CONTROL DE CORTE LASER | DETALLE DE PIEZAS PENDIENTES Y TERMINADAS"
    c1.font = title_font
    c1.fill = title_fill
    c1.alignment = center_align
    ws1.row_dimensions[1].height = 28
    
    tot_plan = int(df_m['pzas_planeadas'].sum()) if not df_m.empty else 0
    tot_cort = int(df_m['pzas_cortadas'].sum()) if not df_m.empty else 0
    tot_pend = int(df_m['pzas_pendientes'].sum()) if not df_m.empty else 0
    
    ws1.merge_cells("A2:K2")
    c2 = ws1["A2"]
    now_str = get_local_now().strftime("%d/%m/%Y %H:%M")
    c2.value = f"OF: {of_number}   |   Fecha: {now_str}   |   Planeadas: {tot_plan:,}   |   Terminadas: {tot_cort:,}   |   Pendientes: {tot_pend:,}"
    c2.font = sub_font
    c2.fill = red_banner
    c2.alignment = center_align
    ws1.row_dimensions[2].height = 20
    ws1.row_dimensions[3].height = 8
    
    cols1 = [
        ("of_number", "OF", 18, center_align),
        ("nido", "Nido", 12, center_align),
        ("no_pieza", "No. Pieza", 20, center_align),
        ("nombre_pieza", "Descripción", 38, left_align),
        ("hojas", "Hojas Req.", 12, center_align),
        ("hojas_cortadas", "Hojas Cortadas", 14, center_align),
        ("cantidad", "Pzas / Hoja", 12, center_align),
        ("pzas_planeadas", "Total Planeadas", 15, center_align),
        ("pzas_cortadas", "Pzas Terminadas", 15, center_align),
        ("pzas_pendientes", "Pzas Pendientes", 15, center_align),
        ("estado", "Estado", 14, center_align),
    ]
    
    ws1.row_dimensions[4].height = 25
    for c_i, (_, name, w, _) in enumerate(cols1, 1):
        cell = ws1.cell(row=4, column=c_i, value=name)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = center_align
        cell.border = thin_border
        ws1.column_dimensions[get_column_letter(c_i)].width = w
        
    curr = 5
    for _, r in df_m.iterrows():
        ws1.row_dimensions[curr].height = 20
        r_fill = alt_fill if curr % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
        for c_i, (k, _, _, al) in enumerate(cols1, 1):
            val = r.get(k, "")
            cell = ws1.cell(row=curr, column=c_i, value=val)
            cell.alignment = al
            cell.border = thin_border
            cell.fill = r_fill
            cell.font = data_font
            
            if k == "pzas_planeadas":
                cell.font = bold_font
            elif k == "pzas_cortadas":
                cell.fill = fill_term
                cell.font = font_term
            elif k == "pzas_pendientes":
                if val > 0:
                    cell.fill = fill_pend
                    cell.font = font_pend
            elif k == "estado":
                if val == "Terminado":
                    cell.fill = fill_ok
                    cell.font = bold_font
        curr += 1
        
    # Totales hoja 1
    ws1.row_dimensions[curr].height = 22
    ws1.cell(row=curr, column=1, value="TOTALES").font = tot_font
    ws1.cell(row=curr, column=1).alignment = center_align
    ws1.cell(row=curr, column=1).fill = tot_fill
    ws1.cell(row=curr, column=1).border = thick_bottom
    
    for c_i in range(2, len(cols1) + 1):
        cell = ws1.cell(row=curr, column=c_i)
        cell.fill = tot_fill
        cell.border = thick_bottom
        k = cols1[c_i - 1][0]
        if k in ["pzas_planeadas", "pzas_cortadas", "pzas_pendientes"]:
            col_let = get_column_letter(c_i)
            cell.value = f"=SUM({col_let}5:{col_let}{curr-1})"
            cell.font = tot_font
            cell.alignment = center_align

    # HOJA 2: RESUMEN POR NÚMERO DE PARTE
    if not df_m.empty:
        ws2 = wb.create_sheet(title="Resumen por Parte")
        
        ws2.merge_cells("A1:H1")
        c1 = ws2["A1"]
        c1.value = f"SIGRAMA — RESUMEN DE CORTE POR NUMERO DE PARTE | {of_number}"
        c1.font = title_font
        c1.fill = title_fill
        c1.alignment = center_align
        ws2.row_dimensions[1].height = 28
        
        ws2.merge_cells("A2:H2")
        c2 = ws2["A2"]
        c2.value = f"Generado: {now_str}   |   Total de Partes Distintas: {df_m['no_pieza'].nunique()}"
        c2.font = sub_font
        c2.fill = red_banner
        c2.alignment = center_align
        ws2.row_dimensions[2].height = 20
        ws2.row_dimensions[3].height = 8
        
        df_res = df_m.groupby(['of_number', 'no_pieza']).agg({
            'nombre_pieza': 'first',
            'pzas_planeadas': 'sum',
            'pzas_cortadas': 'sum',
            'pzas_pendientes': 'sum'
        }).reset_index()
        
        df_res['avance_pct'] = df_res.apply(
            lambda r: (r['pzas_cortadas'] / r['pzas_planeadas']) if r['pzas_planeadas'] > 0 else 0, axis=1
        )
        df_res['estado'] = df_res.apply(
            lambda r: "Terminado" if r['pzas_pendientes'] == 0 else ("En Proceso" if r['pzas_cortadas'] > 0 else "Pendiente"), axis=1
        )
        
        cols2 = [
            ("of_number", "OF", 18, center_align),
            ("no_pieza", "No. Pieza", 22, center_align),
            ("nombre_pieza", "Descripción", 40, left_align),
            ("pzas_planeadas", "Total Planeadas", 15, center_align),
            ("pzas_cortadas", "Total Terminadas", 15, center_align),
            ("pzas_pendientes", "Total Pendientes", 15, center_align),
            ("avance_pct", "% Avance", 12, center_align),
            ("estado", "Estado", 14, center_align)
        ]
        
        ws2.row_dimensions[4].height = 25
        for c_i, (_, name, w, _) in enumerate(cols2, 1):
            cell = ws2.cell(row=4, column=c_i, value=name)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = center_align
            cell.border = thin_border
            ws2.column_dimensions[get_column_letter(c_i)].width = w
            
        curr2 = 5
        for _, r in df_res.iterrows():
            ws2.row_dimensions[curr2].height = 20
            r_fill = alt_fill if curr2 % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
            for c_i, (k, _, _, al) in enumerate(cols2, 1):
                val = r.get(k, "")
                cell = ws2.cell(row=curr2, column=c_i, value=val)
                cell.alignment = al
                cell.border = thin_border
                cell.fill = r_fill
                cell.font = data_font
                
                if k == "pzas_planeadas":
                    cell.font = bold_font
                elif k == "pzas_cortadas":
                    cell.fill = fill_term
                    cell.font = font_term
                elif k == "pzas_pendientes":
                    if val > 0:
                        cell.fill = fill_pend
                        cell.font = font_pend
                elif k == "avance_pct":
                    cell.number_format = '0.0%'
                    cell.font = bold_font
                elif k == "estado":
                    if val == "Terminado":
                        cell.fill = fill_ok
                        cell.font = bold_font
            curr2 += 1
            
        ws2.row_dimensions[curr2].height = 22
        ws2.cell(row=curr2, column=1, value="TOTALES").font = tot_font
        ws2.cell(row=curr2, column=1).alignment = center_align
        ws2.cell(row=curr2, column=1).fill = tot_fill
        ws2.cell(row=curr2, column=1).border = thick_bottom
        
        for c_i in range(2, len(cols2) + 1):
            cell = ws2.cell(row=curr2, column=c_i)
            cell.fill = tot_fill
            cell.border = thick_bottom
            k = cols2[c_i - 1][0]
            if k in ["pzas_planeadas", "pzas_cortadas", "pzas_pendientes"]:
                col_let = get_column_letter(c_i)
                cell.value = f"=SUM({col_let}5:{col_let}{curr2-1})"
                cell.font = tot_font
                cell.alignment = center_align

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def view_avances():
    st.markdown("## 🏭 Panel de Operador - Registro de Avances")
    
    st.markdown(
        f"""
        <div style="background-color: {GRAY}; padding: 15px; border-radius: 8px; margin-bottom: 20px;">
            <h4 style="margin: 0; color: {BLACK};">Filtros de Búsqueda y Estación de Trabajo</h4>
        </div>
        """, unsafe_allow_html=True
    )
    
    # 1. Obtener proyectos y calibres disponibles para los filtros
    conn = get_connection()
    df_proj = pd.read_sql_query("SELECT DISTINCT proyecto FROM ordenes WHERE proyecto IS NOT NULL AND proyecto != ''", conn)
    proyectos_list = ["Todos"] + df_proj["proyecto"].tolist()
    
    # Verificar si calibre existe en la base de datos
    c_info = conn.cursor()
    c_info.execute("PRAGMA table_info(nidos)")
    nidos_cols = [row[1] for row in c_info.fetchall()]
    if "calibre" in nidos_cols:
        df_cal = pd.read_sql_query("SELECT DISTINCT calibre FROM nidos WHERE calibre IS NOT NULL AND calibre != ''", conn)
        calibres_list = ["Todos"] + df_cal["calibre"].tolist()
    else:
        calibres_list = ["Todos"]
    conn.close()

    # Fila 1 de Filtros (3 columnas)
    col_f1, col_f2, col_f3 = st.columns(3)
    
    with col_f1:
        area_seleccionada = st.selectbox("1️⃣ Área / Proceso", PROCESSES, key="avance_area")
        
    with col_f2:
        proyecto_seleccionado = st.selectbox("📂 Proyecto", proyectos_list, key="avance_proyecto_selector")
        
    with col_f3:
        # Filtrar OFs disponibles por el proyecto seleccionado
        conn = get_connection()
        if proyecto_seleccionado == "Todos":
            df_ofs_f = pd.read_sql_query("SELECT DISTINCT of_number FROM ordenes", conn)
        else:
            df_ofs_f = pd.read_sql_query("SELECT DISTINCT of_number FROM ordenes WHERE proyecto = ?", conn, params=(proyecto_seleccionado,))
        conn.close()
        
        todas_ofs = df_ofs_f["of_number"].tolist()
        todas_ofs_opciones = ["Todas"] + todas_ofs
        
        # Sincronizar el selector local con la OF seleccionada en la barra lateral
        default_idx = 0
        global_active_of = st.session_state.get("of_number")
        if global_active_of in todas_ofs_opciones:
            default_idx = todas_ofs_opciones.index(global_active_of)
            
        of_seleccionada = st.selectbox("2️⃣ Orden de Fabricación (OF)", todas_ofs_opciones, index=default_idx, key="avance_of_selector")
        of_number = of_seleccionada

    # Fila 2 de Filtros (3 columnas)
    col_f4, col_f5, col_f6 = st.columns(3)
    
    with col_f4:
        calibre_seleccionado = st.selectbox("📏 Calibre", calibres_list, key="avance_calibre_selector")
        
    with col_f5:
        maquinas_por_area = {
            "Corte": ["Laser 1", "Laser 2", "Laser 3", "Laser 4"],
            "Rebabeo": ["Lijadora 1", "Lijadora 2"],
            "Doblez": ["Dobladora 1", "Dobladora 2", "Dobladora 3", "Dobladora 4"],
            "Barrenado": ["Manual"],
            "Pintura": ["Linea Continua", "Linea de Pintura"],
            "Liberado": ["Manual", "N/A"],
            "Empaque": ["Manual", "N/A"],
            "Ingenieria": ["N/A"]
        }
        lista_maquinas = maquinas_por_area.get(area_seleccionada, ["N/A"])
        maquina = st.selectbox("3️⃣ Máquina", lista_maquinas, key="avance_maquina")
        
    with col_f6:
        ops_list = get_operadores_por_area(area_seleccionada)
        if ops_list:
            operador = st.selectbox("4️⃣ Operador", ops_list, key="avance_operador")
        else:
            operador = st.text_input("4️⃣ Operador (Nombre o Nómina)", key="avance_operador")
        
    # Obtener piezas desde la Base de Datos
    df_todas = get_todas_piezas(of_number)
    
    # Aplicar filtro de calibre en python
    if calibre_seleccionado != "Todos" and "calibre" in df_todas.columns:
        df_todas = df_todas[df_todas['calibre'] == calibre_seleccionado]
        
    if df_todas.empty:
        st.warning("⚠️ No se encontraron piezas que coincidan con los filtros seleccionados.")
        return
        
    nidos_list = df_todas['nido'].unique().tolist()
    
    is_ingenieria = (area_seleccionada == "Ingenieria")
    is_corte = (area_seleccionada == "Corte")
    is_post_corte = not is_ingenieria and not is_corte
    
    nidos_seleccionados = []
    if is_corte:
        # Calcular WIP de Corte: solo contar piezas de nidos AUN NO TERMINADOS
        # Un nido de Corte esta terminado si hojas_cortadas >= hojas_requeridas
        conn = get_connection()
        c_wip = conn.cursor()
        if of_number == "Todas":
            c_wip.execute("SELECT of_number, nido, hojas FROM nidos")
        else:
            c_wip.execute("SELECT of_number, nido, hojas FROM nidos WHERE of_number=?", (of_number,))
        nidos_data = c_wip.fetchall()

        total_wip_corte = 0
        for (of_n, nido_n, hojas_req) in nidos_data:
            hojas_req = int(hojas_req) if hojas_req else 1
            c_wip.execute(
                "SELECT COUNT(DISTINCT hoja) FROM avances WHERE of_number=? AND nido=? AND area='Corte'",
                (of_n, nido_n)
            )
            hojas_cortadas = c_wip.fetchone()[0]
            if hojas_cortadas < hojas_req:   # Nido aun pendiente
                c_wip.execute(
                    "SELECT SUM(p.cantidad * n.hojas) FROM piezas p "
                    "JOIN nidos n ON p.of_number=n.of_number AND p.nido=n.nido "
                    "WHERE p.of_number=? AND p.nido=?",
                    (of_n, nido_n)
                )
                plan_nido = c_wip.fetchone()[0] or 0
                c_wip.execute(
                    "SELECT SUM(cantidad) FROM avances WHERE of_number=? AND nido=? AND area='Corte'",
                    (of_n, nido_n)
                )
                avanzado_nido = c_wip.fetchone()[0] or 0
                total_wip_corte += max(0, plan_nido - avanzado_nido)
        conn.close()
        
        st.markdown(
            f"""
            <div style="background-color: #f8f9fa; border-left: 8px solid #EC2024; padding: 25px; border-radius: 8px; margin-bottom: 20px; box-shadow: 0 4px 6px rgba(0,0,0,0.1);">
                <h3 style="margin: 0; color: #555; font-size: 1.4rem; text-transform: uppercase;">Sumatoria WIP - Corte (Piezas por cortar)</h3>
                <div style="display: flex; align-items: baseline; gap: 10px;">
                    <h1 style="margin: 0; color: #EC2024; font-size: 4.5rem; font-weight: 900; line-height: 1;">{int(total_wip_corte)}</h1>
                    <span style="font-size: 1.8rem; font-weight: bold; color: #666;">piezas disponibles para procesar</span>
                </div>
            </div>
            """, unsafe_allow_html=True
        )
        
        col_corte_dl, _ = st.columns([1, 1])
        with col_corte_dl:
            excel_corte_bytes = generate_corte_detalle_excel(of_number)
            clean_of_c = re.sub(r'[^a-zA-Z0-9_-]', '_', of_number)
            fecha_c_tag = get_local_now().strftime("%Y%m%d_%H%M")
            st.download_button(
                label="📥 Descargar Detalle de Piezas en Corte (.xlsx)",
                data=excel_corte_bytes,
                file_name=f"Detalle_Corte_{clean_of_c}_{fecha_c_tag}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                type="secondary",
                key="btn_dl_corte_piezas_excel"
            )

        # Toggle para cambiar a modo reposición
        modo_reposicion = st.checkbox("🔄 Registrar Reposiciones / Re-cuts (Piezas sueltas re-cortadas por Scrap)", key="corte_modo_reposicion")
        
        if modo_reposicion:
            st.markdown("### 📋 Registro de Reposición de Piezas en Corte")
            st.markdown("👇 **Ingresa la cantidad de piezas que re-cortaste para reponer Scrap:**")
            
            # Agrupar piezas por número de pieza
            df_plan_piezas = df_todas.copy()
            df_plan_piezas['hojas'] = pd.to_numeric(df_plan_piezas['hojas'], errors='coerce').fillna(1).astype(int)
            df_plan_piezas['cantidad'] = pd.to_numeric(df_plan_piezas['cantidad'], errors='coerce').fillna(0).astype(int)
            df_plan_piezas['req_total'] = df_plan_piezas['cantidad'] * df_plan_piezas['hojas']
            
            df_piezas_corte = df_plan_piezas.groupby(['no_pieza']).agg({
                'nombre_pieza': 'first',
                'req_total': 'sum'
            }).reset_index()
            
            # Consultar scrap total (todas las áreas) y reposiciones ya hechas
            conn = get_connection()
            if of_number == "Todas":
                # IMPORTANTE: incluir of_number en GROUP BY para no mezclar OFs con mismo no_pieza
                df_re_c = pd.read_sql_query("SELECT of_number, no_pieza, SUM(cantidad) as rechazadas FROM rechazos GROUP BY of_number, no_pieza", conn)
                df_rep_c = pd.read_sql_query("SELECT of_number, no_pieza, SUM(cantidad) as repuestas FROM avances WHERE area='Corte' AND nido='REPOSICION' GROUP BY of_number, no_pieza", conn)
            else:
                df_re_c = pd.read_sql_query("SELECT no_pieza, SUM(cantidad) as rechazadas FROM rechazos WHERE of_number=? GROUP BY no_pieza", conn, params=(of_number,))
                df_rep_c = pd.read_sql_query("SELECT no_pieza, SUM(cantidad) as repuestas FROM avances WHERE of_number=? AND area='Corte' AND nido='REPOSICION' GROUP BY no_pieza", conn, params=(of_number,))
            conn.close()
            
            df_piezas_corte = df_piezas_corte.merge(df_re_c, on='no_pieza', how='left')
            df_piezas_corte = df_piezas_corte.merge(df_rep_c, on='no_pieza', how='left')
            df_piezas_corte['rechazadas'] = df_piezas_corte['rechazadas'].fillna(0).astype(int)
            df_piezas_corte['repuestas'] = df_piezas_corte['repuestas'].fillna(0).astype(int)
            
            # Pendientes de reponer = Rechazadas - Repuestas
            df_piezas_corte['Pendientes de Reponer'] = df_piezas_corte['rechazadas'] - df_piezas_corte['repuestas']
            df_piezas_corte['Pendientes de Reponer'] = df_piezas_corte['Pendientes de Reponer'].apply(lambda x: max(0, x))
            
            # Filtrar: solo mostrar piezas que tienen rechazos y que aún tienen pendientes de reponer
            df_piezas_corte = df_piezas_corte[(df_piezas_corte['rechazadas'] > 0) & (df_piezas_corte['Pendientes de Reponer'] > 0)].copy()
            
            if df_piezas_corte.empty:
                st.success("✅ ¡Excelente! No hay piezas con rechazos (Scrap) pendientes de reponer para esta OF.")
            else:
                df_piezas_corte['Re-cuts / Reposición'] = 0
                
                df_rep_edit = st.data_editor(
                    df_piezas_corte[['no_pieza', 'nombre_pieza', 'req_total', 'rechazadas', 'repuestas', 'Pendientes de Reponer', 'Re-cuts / Reposición']],
                    column_config={
                        "no_pieza": st.column_config.TextColumn("No. Parte", disabled=True),
                        "nombre_pieza": st.column_config.TextColumn("Descripción", disabled=True),
                        "req_total": st.column_config.NumberColumn("Planeadas", disabled=True),
                        "rechazadas": st.column_config.NumberColumn("Rechazadas (Scrap)", disabled=True),
                        "repuestas": st.column_config.NumberColumn("Repuestas Previamente", disabled=True),
                        "Pendientes de Reponer": st.column_config.NumberColumn("Pendientes de Reponer", disabled=True),
                        "Re-cuts / Reposición": st.column_config.NumberColumn("Cantidad a Reponer", min_value=0, step=1, disabled=False)
                    },
                    use_container_width=True,
                    hide_index=True,
                    key="editor_corte_reposiciones"
                )
                
                if st.button("✅ Registrar Reposiciones en Corte", type="primary", use_container_width=True):
                    if not operador.strip():
                        st.error("⚠️ Por favor selecciona un Operador válido antes de guardar.")
                        st.stop()
                        
                    df_to_save = df_rep_edit[df_rep_edit['Re-cuts / Reposición'] > 0].copy()
                    if df_to_save.empty:
                        st.warning("⚠️ No has capturado ninguna cantidad de reposición.")
                    else:
                        # Validar el límite máximo
                        sobrepasadas = df_to_save[df_to_save['Re-cuts / Reposición'] > df_to_save['Pendientes de Reponer']]
                        if not sobrepasadas.empty:
                            st.error("❌ No puedes reponer más piezas de las que fueron rechazadas. Revisa las cantidades capturadas.")
                            st.stop()
                            
                        df_to_save['Terminadas'] = df_to_save['Re-cuts / Reposición']
                        # Guardar avance en Corte con nido='REPOSICION' y hoja=None
                        save_avances_mixto(of_number, "REPOSICION", "Corte", False, df_to_save, None, operador, maquina, None)
                        st.success("🎉 ¡Reposiciones registradas exitosamente en Corte!")
                        st.rerun()
        else:
            st.markdown("👇 **Selecciona un Nido de la tabla haciendo clic en la fila correspondiente:**")
            
            # Estado de cada nido: TERMINADO, EN PROCESO, o PENDIENTE según hojas cortadas (para Corte)
            conn = get_connection()
            c = conn.cursor()
            if of_number == "Todas":
                c.execute("SELECT n.of_number, n.nido, n.hojas, COUNT(DISTINCT a.hoja) as cortadas "
                          "FROM nidos n LEFT JOIN avances a ON n.of_number=a.of_number AND n.nido=a.nido AND a.area='Corte' "
                          "GROUP BY n.of_number, n.nido, n.hojas")
                nido_status = {}
                for r in c.fetchall():
                    of_n, nido_n, hojas, cortadas = r[0], r[1], int(r[2]) if r[2] else 1, int(r[3]) if r[3] else 0
                    if cortadas >= hojas:
                        nido_status[f"{of_n} | {nido_n}"] = f"✅ Terminado ({cortadas}/{hojas})"
                    elif cortadas > 0:
                        nido_status[f"{of_n} | {nido_n}"] = f"🔄 En Proceso ({cortadas}/{hojas})"
                    else:
                        nido_status[f"{of_n} | {nido_n}"] = f"⏳ Pendiente (0/{hojas})"
                
                nidos_list_unique = (df_todas['of_number'] + ' | ' + df_todas['nido']).unique().tolist()
                df_nidos_list = pd.DataFrame({"Nido": nidos_list_unique})
                df_nidos_list["Estado"] = df_nidos_list["Nido"].apply(
                    lambda x: nido_status.get(x, "⏳ Pendiente (0/1)")
                )
            else:
                c.execute("SELECT n.nido, n.hojas, COUNT(DISTINCT a.hoja) as cortadas "
                          "FROM nidos n LEFT JOIN avances a ON n.of_number=a.of_number AND n.nido=a.nido AND a.area='Corte' "
                          "WHERE n.of_number=? GROUP BY n.nido, n.hojas", (of_number,))
                nido_status = {}
                for r in c.fetchall():
                    nido_n, hojas, cortadas = r[0], int(r[1]) if r[1] else 1, int(r[2]) if r[2] else 0
                    if cortadas >= hojas:
                        nido_status[nido_n] = f"✅ Terminado ({cortadas}/{hojas})"
                    elif cortadas > 0:
                        nido_status[nido_n] = f"🔄 En Proceso ({cortadas}/{hojas})"
                    else:
                        nido_status[nido_n] = f"⏳ Pendiente (0/{hojas})"
                
                df_nidos_list = pd.DataFrame({"Nido": nidos_list})
                df_nidos_list["Estado"] = df_nidos_list["Nido"].apply(
                    lambda x: nido_status.get(x, "⏳ Pendiente (0/1)")
                )
            conn.close()
            
            event = st.dataframe(
                df_nidos_list,
                use_container_width=True,
                hide_index=True,
                on_select="rerun",
                selection_mode="multi-row",
                key="nido_selector",
                height=200
            )
            
            selected_rows = event.selection.rows
            if selected_rows:
                for idx in selected_rows:
                    if idx < len(df_nidos_list):
                        nidos_seleccionados.append(df_nidos_list.iloc[idx]['Nido'])

    if is_ingenieria:
        # Calcular total de piezas requeridas por número de parte
        df_plan_piezas = df_todas.copy()
        df_plan_piezas['hojas'] = pd.to_numeric(df_plan_piezas['hojas'], errors='coerce').fillna(1).astype(int)
        df_plan_piezas['cantidad'] = pd.to_numeric(df_plan_piezas['cantidad'], errors='coerce').fillna(0).astype(int)
        df_plan_piezas['total_req'] = df_plan_piezas['cantidad'] * df_plan_piezas['hojas']
        df_sum_req = df_plan_piezas.groupby(['of_number', 'no_pieza'])['total_req'].sum().reset_index()
        
        # Unique parts
        df_partes = df_todas.drop_duplicates(subset=['of_number', 'no_pieza']).copy()
        
        # Check which ones are already completed in Ingenieria
        conn = get_connection()
        c = conn.cursor()
        if of_number == "Todas":
            c.execute("SELECT DISTINCT of_number || '-' || no_pieza FROM avances WHERE area = 'Ingenieria'")
        else:
            c.execute("SELECT DISTINCT of_number || '-' || no_pieza FROM avances WHERE of_number = ? AND area = 'Ingenieria'", (of_number,))
            
        terminadas = [row[0] for row in c.fetchall()]
        conn.close()
        
        # Create df_edit
        df_edit = df_partes[['of_number', 'no_pieza', 'nombre_pieza']].copy()
        df_edit['Diseñada'] = df_edit.apply(lambda row: f"{row['of_number']}-{row['no_pieza']}" in terminadas, axis=1)
        
        # Unir para obtener la cantidad planeada
        df_edit = df_edit.merge(df_sum_req, on=['of_number', 'no_pieza'], how='left')
        df_edit['total_req'] = df_edit['total_req'].fillna(0).astype(int)
        
        # Calcular WIP de Ingeniería: contar partes únicas pendientes de diseñar
        df_pendientes = df_edit[df_edit['Diseñada'] == False].copy()
        total_wip_ingenieria = len(df_pendientes)   # Número de partes únicas, no piezas físicas
        total_partes = len(df_edit)
        
        # Mostrar tarjeta WIP
        st.markdown(
            f"""
            <div style="background-color: #f8f9fa; border-left: 8px solid #EC2024; padding: 25px; border-radius: 8px; margin-bottom: 20px; box-shadow: 0 4px 6px rgba(0,0,0,0.1);">
                <h3 style="margin: 0; color: #555; font-size: 1.4rem; text-transform: uppercase;">WIP Ingeniería — Números de Parte por Diseñar</h3>
                <div style="display: flex; align-items: baseline; gap: 10px;">
                    <h1 style="margin: 0; color: #EC2024; font-size: 4.5rem; font-weight: 900; line-height: 1;">{total_wip_ingenieria}</h1>
                    <span style="font-size: 1.8rem; font-weight: bold; color: #666;">/ {total_partes} partes pendientes de diseñar</span>
                </div>
            </div>
            """, unsafe_allow_html=True
        )
        
        # Mostrar tabla de pendientes de Diseñar
        if not df_pendientes.empty:
            df_pend_show = df_pendientes.rename(columns={'no_pieza': 'NP', 'total_req': 'CANTIDAD'})
            st.markdown("### ⏳ Piezas Pendientes de Diseñar")
            st.dataframe(df_pend_show[['NP', 'CANTIDAD']], use_container_width=True, hide_index=True)
            
        st.markdown(f"### 📋 Números de Parte de la OF: **{of_number}**")
        st.markdown("👇 **Para INGENIERÍA, marca las partes que ya has terminado de diseñar/programar:**")
        
        edited_df = st.data_editor(
            df_edit,
            use_container_width=True,
            hide_index=True,
            column_config={
                "of_number": st.column_config.TextColumn("OF", disabled=True),
                "no_pieza": st.column_config.TextColumn("No. Parte (NP)", disabled=True),
                "nombre_pieza": st.column_config.TextColumn("Descripción", disabled=True),
                "total_req": st.column_config.NumberColumn("Cantidad Planeada", disabled=True),
                "Diseñada": st.column_config.CheckboxColumn("¿Terminada?")
            },
            height=250
        )
        
        if st.button("✅ Guardar Avances de Ingeniería", type="primary"):
            try:
                # Construir clave compuesta OF-no_pieza para comparar correctamente
                edited_df["_key"] = edited_df["of_number"].astype(str) + "-" + edited_df["no_pieza"].astype(str)
                nuevas_terminadas = edited_df[
                    edited_df["Diseñada"] & (~edited_df["_key"].isin(terminadas))
                ].copy()
                
                if nuevas_terminadas.empty:
                    st.info("No hay partes nuevas por registrar. Todas las partes marcadas ya estaban guardadas.")
                else:
                    nuevas_terminadas["cantidad"] = 1
                    nuevas_terminadas["Terminadas"] = 1
                    save_avances_mixto(of_number, "N/A", area_seleccionada, False, nuevas_terminadas, None, operador, maquina, None)
                    st.success(f"🎉 ¡{len(nuevas_terminadas)} partes marcadas como terminadas en Ingeniería!")
                    st.rerun()
            except Exception as e:
                st.error(f"❌ Error al guardar avances de Ingeniería: {e}")

                
    elif is_corte and nidos_seleccionados:
        if len(nidos_seleccionados) > 1:
            st.markdown(f"### 📦 Registro Masivo de Avances: **{len(nidos_seleccionados)} Nidos Seleccionados**")
            st.info(f"👉 **Nidos seleccionados para completar:** {', '.join(nidos_seleccionados)}")
            
            # Botón de registro masivo
            if st.button(f"⏩ Registrar TODAS las Hojas de los {len(nidos_seleccionados)} Nidos Seleccionados", type="primary", use_container_width=True):
                if not operador.strip():
                    st.error("⚠️ Por favor selecciona un Operador válido antes de guardar.")
                    st.stop()
                    
                with st.spinner("Registrando avances de nidos seleccionados..."):
                    total_nidos_completados = 0
                    total_hojas_completadas = 0
                    for nido_item in nidos_seleccionados:
                        actual_of = of_number
                        actual_nido = nido_item
                        if " | " in nido_item:
                            actual_of, actual_nido = nido_item.split(" | ", 1)
                            
                        # Verificar si este nido ya se completó en esta área
                        df_avances = get_avances_nido(actual_of, actual_nido)
                        areas_terminadas = df_avances['area'].tolist() if not df_avances.empty else []
                        if "Corte" in areas_terminadas:
                            continue
                            
                        # Obtener piezas del nido específico
                        df_nido = df_todas[(df_todas['nido'] == actual_nido) & (df_todas['of_number'] == actual_of)].copy()
                        if df_nido.empty:
                            continue
                        total_hojas = int(df_nido['hojas'].iloc[0]) if 'hojas' in df_nido.columns else 1
                        
                        # Consultar hojas ya cortadas
                        conn = get_connection()
                        c = conn.cursor()
                        c.execute("SELECT DISTINCT hoja FROM avances WHERE of_number=? AND nido=? AND area='Corte' AND hoja IS NOT NULL", (actual_of, actual_nido))
                        hojas_cortadas = {int(row[0]) for row in c.fetchall()}
                        conn.close()
                        
                        nido_changed = False
                        # Registrar todas las hojas faltantes para este nido
                        for h in range(1, total_hojas + 1):
                            if h not in hojas_cortadas:
                                df_terminadas = df_nido[['no_pieza', 'nombre_pieza', 'cantidad']].copy()
                                df_terminadas["Terminadas"] = df_terminadas["cantidad"]
                                df_terminadas["of_number"] = actual_of
                                
                                save_avances_mixto(actual_of, actual_nido, area_seleccionada, is_corte, df_terminadas, None, operador, maquina, h)
                                total_hojas_completadas += 1
                                nido_changed = True
                        if nido_changed or len(hojas_cortadas) >= total_hojas:
                            total_nidos_completados += 1
                            
                st.success(f"🎉 ¡Sincronización Masiva Completa! Se registraron {total_nidos_completados} nidos y {total_hojas_completadas} hojas exitosamente en Corte.")
                st.rerun()
        else:
            # Caso de un solo nido seleccionado (Lógica tradicional paso a paso)
            nido_seleccionado = nidos_seleccionados[0]
            st.markdown(f"### 📋 Detalles del Nido: **{nido_seleccionado}**")
            
            actual_of = of_number
            actual_nido = nido_seleccionado
            if " | " in nido_seleccionado:
                actual_of, actual_nido = nido_seleccionado.split(" | ", 1)
                
            # Verificar si este nido ya se completó en esta área
            df_avances = get_avances_nido(actual_of, actual_nido)
            areas_terminadas = df_avances['area'].tolist() if not df_avances.empty else []
            
            if area_seleccionada in areas_terminadas:
                st.success(f"✅ Este nido ya fue marcado como TERMINADO en **{area_seleccionada}**.")
                st.stop() # No permitir registrar de nuevo
                
            # Obtener piezas del nido específico
            df_nido = df_todas[(df_todas['nido'] == actual_nido) & (df_todas['of_number'] == actual_of)].copy()
            total_hojas = int(df_nido['hojas'].iloc[0]) if 'hojas' in df_nido.columns and not df_nido.empty else 1
            
            # Consultar en DB qué hojas ya fueron cortadas
            conn = get_connection()
            c = conn.cursor()
            c.execute("SELECT DISTINCT hoja FROM avances WHERE of_number=? AND nido=? AND area='Corte' AND hoja IS NOT NULL", (actual_of, actual_nido))
            hojas_cortadas = {int(row[0]) for row in c.fetchall()}
            conn.close()
            
            # Encontrar la primera hoja pendiente (1-indexed)
            hoja_actual = 1
            for h in range(1, total_hojas + 1):
                if h not in hojas_cortadas:
                    hoja_actual = h
                    break
            else:
                hoja_actual = total_hojas + 1
            
            if hoja_actual > total_hojas:
                st.success(f"✅ Todas las hojas ({total_hojas}/{total_hojas}) de este Nesteo ya fueron cortadas.")
            else:
                st.markdown(f"👇 **CORTE: Registrando Hoja {hoja_actual} de {total_hojas}**")
                st.markdown(f"Las cantidades mostradas abajo corresponden **solamente a las piezas que salen de esta hoja**.")
                
                df_edit = df_nido[['no_pieza', 'nombre_pieza', 'cantidad']].copy()
                df_edit['Rechazos'] = 0
                df_edit['Motivo'] = ""
                
                edited_df = st.data_editor(
                    df_edit,
                    use_container_width=True,
                    hide_index=True,
                    column_config={
                        "no_pieza": st.column_config.TextColumn("No. Pieza", disabled=True),
                        "nombre_pieza": st.column_config.TextColumn("Descripción", disabled=True),
                        "cantidad": st.column_config.NumberColumn("Cantidad (x Hoja)", disabled=True),
                        "Rechazos": st.column_config.NumberColumn("Cant. Rechazada", min_value=0, step=1),
                        "Motivo": st.column_config.TextColumn("Motivo de rechazo")
                    },
                    height=200
                )
                
                col_b1, col_b2 = st.columns(2)
                with col_b1:
                    if st.button(f"✅ Registrar Hoja {hoja_actual} Terminada", type="primary", use_container_width=True):
                        if not operador.strip():
                            st.error("⚠️ Por favor selecciona un Operador válido.")
                            st.stop()
                            
                        df_terminadas = edited_df.copy()
                        df_terminadas["Terminadas"] = df_terminadas["cantidad"] # Avanza exactamente lo de 1 hoja
                        df_terminadas["of_number"] = actual_of
                        
                        df_rechazos = edited_df[edited_df["Rechazos"] > 0].copy()
                        if not df_rechazos.empty:
                            df_rechazos["of_number"] = actual_of
                            
                        save_avances_mixto(actual_of, actual_nido, area_seleccionada, is_corte, df_terminadas, df_rechazos, operador, maquina, hoja_actual)
                        st.success(f"🎉 ¡Hoja {hoja_actual} registrada en Corte!")
                        st.rerun()
                
                with col_b2:
                    if st.button(f"⏩ Registrar TODAS las Hojas ({hoja_actual} a {total_hojas})", type="secondary", use_container_width=True):
                        if not operador.strip():
                            st.error("⚠️ Por favor selecciona un Operador válido.")
                            st.stop()
                            
                        with st.spinner("Registrando todas las hojas..."):
                            registered_count = 0
                            for h in range(hoja_actual, total_hojas + 1):
                                if h not in hojas_cortadas:
                                    df_terminadas = edited_df.copy()
                                    df_terminadas["Terminadas"] = df_terminadas["cantidad"]
                                    df_terminadas["of_number"] = actual_of
                                    
                                    df_rechazos = edited_df[edited_df["Rechazos"] > 0].copy()
                                    if not df_rechazos.empty:
                                        df_rechazos["of_number"] = actual_of
                                        
                                    save_avances_mixto(actual_of, actual_nido, area_seleccionada, is_corte, df_terminadas, df_rechazos, operador, maquina, h)
                                    registered_count += 1
                        st.success(f"🎉 ¡{registered_count} Hojas registradas exitosamente en Corte!")
                        st.rerun()


    elif is_post_corte:
        st.markdown(f"### 📋 Números de Parte de la OF: **{of_number}**")
        st.markdown(f"👇 **Para {area_seleccionada}, registra las piezas que vayas terminando.**")
        
        # 1. Filtrar piezas que SÍ pasan por esta área
        df_piezas = df_todas.copy()
        if 'hojas' not in df_piezas.columns:
            df_piezas['hojas'] = 1
        df_piezas['total_requeridas'] = df_piezas['cantidad'] * df_piezas['hojas']
        
        # Verificar si la pieza pasa por area_seleccionada y obtener su area anterior
        def get_area_anterior(ruta_str, current_area):
            procesos = [p.strip() for p in str(ruta_str).split(',') if p.strip()]
            if current_area in procesos:
                idx = procesos.index(current_area)
                if idx > 0:
                    return procesos[idx - 1]
            return None
            
        df_piezas['pasa_por_aqui'] = df_piezas['ruta'].apply(lambda x: area_seleccionada in [p.strip() for p in str(x).split(',') if p.strip()])
        df_piezas = df_piezas[df_piezas['pasa_por_aqui']]
        
        if df_piezas.empty:
            st.warning(f"⚠️ No hay piezas programadas para pasar por {area_seleccionada} en esta OF.")
            st.stop()
            
        df_piezas['area_anterior'] = df_piezas['ruta'].apply(lambda x: get_area_anterior(x, area_seleccionada))
        
        # Agrupar por pieza (sumando total requeridas)
        df_agrupado = df_piezas.groupby(['of_number', 'no_pieza']).agg({
            'nombre_pieza': 'first',
            'total_requeridas': 'sum',
            'area_anterior': 'first' # Asumimos misma ruta para misma pieza
        }).reset_index()
        
        # 2. Cargar todos los avances y rechazos de la OF de una sola vez
        conn = get_connection()
        if of_number == "Todas":
            df_avances_all = pd.read_sql_query("SELECT of_number, no_pieza, area, SUM(cantidad) as cantidad FROM avances GROUP BY of_number, no_pieza, area", conn)
            df_rechazos_all = pd.read_sql_query("SELECT of_number, no_pieza, area, SUM(cantidad) as cantidad FROM rechazos GROUP BY of_number, no_pieza, area", conn)
        else:
            df_avances_all = pd.read_sql_query("SELECT of_number, no_pieza, area, SUM(cantidad) as cantidad FROM avances WHERE of_number=? GROUP BY of_number, no_pieza, area", conn, params=(of_number,))
            df_rechazos_all = pd.read_sql_query("SELECT of_number, no_pieza, area, SUM(cantidad) as cantidad FROM rechazos WHERE of_number=? GROUP BY of_number, no_pieza, area", conn, params=(of_number,))
        conn.close()
        
        # 3. Calcular métricas por pieza
        def get_wip(row):
            area_ant = row['area_anterior']
            if pd.isna(area_ant) or not area_ant:
                return 0 # Ej: Si es el primer proceso y no es corte
            wip = df_avances_all[(df_avances_all['no_pieza'] == row['no_pieza']) & (df_avances_all['of_number'] == row['of_number']) & (df_avances_all['area'] == area_ant)]['cantidad'].sum()
            return int(wip)
            
        def get_terminadas_ant(row):
            term = df_avances_all[(df_avances_all['no_pieza'] == row['no_pieza']) & (df_avances_all['of_number'] == row['of_number']) & (df_avances_all['area'] == area_seleccionada)]['cantidad'].sum()
            return int(term)
            
        def get_rechazadas_ant(row):
            rech = df_rechazos_all[(df_rechazos_all['no_pieza'] == row['no_pieza']) & (df_rechazos_all['of_number'] == row['of_number']) & (df_rechazos_all['area'] == area_seleccionada)]['cantidad'].sum()
            return int(rech)
            
        df_agrupado['wip'] = df_agrupado.apply(get_wip, axis=1)
        df_agrupado['terminadas_ant'] = df_agrupado.apply(get_terminadas_ant, axis=1)
        df_agrupado['rechazadas_ant'] = df_agrupado.apply(get_rechazadas_ant, axis=1)
        
        df_edit = df_agrupado.copy()
        
        df_edit['Pendiente Disponible'] = df_edit['wip'] - df_edit['terminadas_ant'] - df_edit['rechazadas_ant']
        df_edit['Pendiente Total OF'] = df_edit['total_requeridas'] - df_edit['terminadas_ant'] - df_edit['rechazadas_ant']
        
        # Evitar negativos
        df_edit['Pendiente Disponible'] = df_edit['Pendiente Disponible'].apply(lambda x: max(0, x))
        df_edit['Pendiente Total OF'] = df_edit['Pendiente Total OF'].apply(lambda x: max(0, x))
        
        # Filtrar solo las que tienen Pendiente Total y además tienen WIP Disponible o avance en esta área
        mask_pendiente = df_edit['Pendiente Total OF'] > 0
        mask_wip = (df_edit['Pendiente Disponible'] > 0) | (df_edit['terminadas_ant'] > 0) | (df_edit['rechazadas_ant'] > 0)
        df_edit = df_edit[mask_pendiente & mask_wip].copy()
        
        if df_edit.empty:
            st.success(f"✅ ¡Todas las piezas de la OF ya fueron procesadas en {area_seleccionada}!")
        else:
            df_edit['Terminadas'] = 0
            df_edit['Rechazos'] = 0
            df_edit['Motivo'] = ""
            df_edit['✅ Todo a Buenas'] = False
            
            total_wip_actual = df_edit['Pendiente Disponible'].sum()
            
            st.markdown(
                f"""
                <div style="background-color: #f8f9fa; border-left: 8px solid #EC2024; padding: 25px; border-radius: 8px; margin-bottom: 20px; box-shadow: 0 4px 6px rgba(0,0,0,0.1);">
                    <h3 style="margin: 0; color: #555; font-size: 1.4rem; text-transform: uppercase;">Sumatoria WIP - {area_seleccionada}</h3>
                    <div style="display: flex; align-items: baseline; gap: 10px;">
                        <h1 style="margin: 0; color: #EC2024; font-size: 4.5rem; font-weight: 900; line-height: 1;">{int(total_wip_actual)}</h1>
                        <span style="font-size: 1.8rem; font-weight: bold; color: #666;">piezas disponibles para procesar</span>
                    </div>
                </div>
                """, unsafe_allow_html=True
            )
            
            # Reordenar columnas para cumplir con la solicitud
            col_order = [
                "✅ Todo a Buenas", "of_number", "no_pieza", "nombre_pieza", "total_requeridas", "Pendiente Disponible", 
                "terminadas_ant", "rechazadas_ant", 
                "Terminadas", "Rechazos", "Motivo", "Pendiente Total OF"
            ]
            df_edit = df_edit[col_order]
            
            gb = GridOptionsBuilder.from_dataframe(df_edit)
            
            # Configurar columnas no editables
            gb.configure_column("✅ Todo a Buenas", header_name="✅ Completar", editable=True, width=120)
            gb.configure_column("of_number", header_name="OF", width=100, pinned='left')
            gb.configure_column("no_pieza", header_name="No. Pieza", width=120)
            gb.configure_column("nombre_pieza", header_name="Descripción", width=250)
            gb.configure_column("total_requeridas", header_name="Totales (OF)", width=100, cellStyle={'backgroundColor': 'black', 'color': 'white'})
            gb.configure_column("Pendiente Disponible", header_name="WIP Disp.", width=100, cellStyle={'backgroundColor': '#d4edda', 'color': 'black'})
            gb.configure_column("terminadas_ant", header_name="Proc. Ant.", width=100, cellStyle={'color': 'blue'})
            gb.configure_column("rechazadas_ant", header_name="Rech. Ant.", width=100, cellStyle={'color': 'red'})
            gb.configure_column("Pendiente Total OF", header_name="Pend. Total", width=110, cellStyle={'color': 'red', 'fontWeight': 'bold', 'fontSize': '120%'})
            
            # Configurar columnas editables
            gb.configure_column("Terminadas", header_name="Proc. Buenas", editable=True, type=["numericColumn","numberColumnFilter"], cellStyle={'backgroundColor': '#d1ecf1', 'fontWeight': 'bold', 'fontSize': '120%', 'color': 'black'})
            gb.configure_column("Rechazos", header_name="Proc. Malas", editable=True, type=["numericColumn","numberColumnFilter"], cellStyle={'backgroundColor': '#f8d7da', 'fontWeight': 'bold', 'fontSize': '120%', 'color': 'black'})
            gb.configure_column("Motivo", header_name="Motivo de rechazo", editable=True, width=150)
            
            # Configurar Javascript para copiar automáticamente cuando se presiona el Checkbox
            on_cell_val_changed = JsCode("""
            function(event) {
                if (event.column.colId === '✅ Todo a Buenas') {
                    if (event.newValue === true) {
                        event.node.setDataValue('Terminadas', event.data['Pendiente Disponible']);
                    } else {
                        event.node.setDataValue('Terminadas', 0);
                    }
                }
            }
            """)
            gb.configure_grid_options(onCellValueChanged=on_cell_val_changed)
            
            grid_options = gb.build()
            
            grid_response = AgGrid(
                df_edit,
                gridOptions=grid_options,
                update_mode=GridUpdateMode.MODEL_CHANGED,
                data_return_mode=DataReturnMode.FILTERED_AND_SORTED,
                fit_columns_on_grid_load=True,
                allow_unsafe_jscode=True,
                theme='streamlit',
                height=250
            )
            
            edited_df = pd.DataFrame(grid_response['data'])
            
            # Cast editables to numeric just in case AgGrid returns them as strings
            edited_df["Terminadas"] = pd.to_numeric(edited_df["Terminadas"], errors='coerce').fillna(0).astype(int)
            edited_df["Rechazos"] = pd.to_numeric(edited_df["Rechazos"], errors='coerce').fillna(0).astype(int)
            
            col_b1, col_b2 = st.columns([1, 1])
            with col_b1:
                btn_avance = st.button(f"✅ Registrar Avance en {area_seleccionada}", type="primary", use_container_width=True)
            with col_b2:
                excel_bytes = generate_wip_table_excel(edited_df, of_number, area_seleccionada, total_wip_actual)
                clean_of = re.sub(r'[^a-zA-Z0-9_-]', '_', of_number)
                fecha_tag = get_local_now().strftime("%Y%m%d_%H%M")
                st.download_button(
                    label=f"📥 Descargar Tabla WIP ({area_seleccionada}) en Excel",
                    data=excel_bytes,
                    file_name=f"WIP_{area_seleccionada}_{clean_of}_{fecha_tag}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                    type="secondary"
                )

            if btn_avance:
                if not operador.strip():
                    st.error("⚠️ Por favor selecciona un Operador válido.")
                    st.stop()
                    
                df_terminadas = edited_df[edited_df["Terminadas"] > 0]
                df_rechazos = edited_df[edited_df["Rechazos"] > 0]
                
                # Validar que no registren más de las pendientes
                if not df_terminadas.empty:
                    sobrepasadas = df_terminadas[df_terminadas["Terminadas"] > df_terminadas["Pendiente Disponible"]]
                    if not sobrepasadas.empty:
                        st.error("❌ No puedes registrar más piezas terminadas de las que tienes 'Disponibles en WIP'.")
                        st.stop()
                
                if df_terminadas.empty and df_rechazos.empty:
                    st.warning("Debes capturar al menos 1 pieza terminada o rechazada.")
                else:
                    # Guardar con nido="N/A"
                    save_avances_mixto(of_number, "N/A", area_seleccionada, False, df_terminadas, df_rechazos, operador, maquina, None)
                    st.success(f"🎉 ¡Avance registrado exitosamente en {area_seleccionada}!")
                    st.rerun()

    # Mostrar historial de movimientos de esta área en la parte inferior (para todas las áreas)
    df_movimientos = get_movimientos_area(of_number, area_seleccionada)
    if not df_movimientos.empty:
        st.markdown("---")
        st.markdown(f"### 🕒 Historial de Movimientos ({area_seleccionada})")
        st.markdown("👇 *Hora y último registro primero. Si necesitas corregir/eliminar un movimiento, dirígete a la sección 3.4 (Correcciones).*")
        
        # Reordenar columnas para poner Fecha (timestamp) primero
        cols_order = ["Fecha", "Tipo", "No. Pieza", "Cantidad", "Operador", "Máquina", "OF", "Motivo"]
        df_movimientos = df_movimientos[[c for c in cols_order if c in df_movimientos.columns]]
        st.dataframe(df_movimientos, use_container_width=True, hide_index=True, height=250)
    else:
        # Mostrar historial de rechazos globales de la OF si no hay movimientos en esta área
        df_rechazos_hist = get_total_rechazos(of_number)
        if not df_rechazos_hist.empty:
            st.markdown("---")
            st.markdown(f"### ⚠️ Historial de Rechazos Globales (OF {of_number})")
            st.dataframe(df_rechazos_hist, use_container_width=True, hide_index=True, height=200)
