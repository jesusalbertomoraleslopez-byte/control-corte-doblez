import fitz
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os
import re
import io
import datetime
import pandas as pd

def procesar_of_pronest(folder_path=None, pdf_resumen_bytes=None, pdf_nido_bytes=None, pdf_pieza_bytes=None, params_orden=None, pdf_bytes_dict=None):
    """
    Procesa los 3 reportes PDF de ProNest y genera los datos estructurados para el Excel (.xlsx)
    requerido por el sistema de Corte y Doblez.
    """
    if pdf_bytes_dict:
        pdf_resumen_bytes = pdf_bytes_dict.get("resumen")
        pdf_nido_bytes = pdf_bytes_dict.get("nido")
        pdf_pieza_bytes = pdf_bytes_dict.get("pieza")

    if folder_path:
        if not os.path.exists(folder_path):
            raise ValueError(f"La ruta de carpeta no existe: {folder_path}")
            
        resumen_path = None
        nido_path = None
        pieza_path = None
        for f in os.listdir(folder_path):
            fl = f.lower()
            if fl.endswith(".pdf"):
                if "resumen de trabajo" in fl or "resumen del trabajo" in fl:
                    resumen_path = os.path.join(folder_path, f)
                elif "detalle del nido de cabezal" in fl or "detalle de nido" in fl:
                    nido_path = os.path.join(folder_path, f)
                elif "detalle de la pieza" in fl or "detalle de pieza" in fl:
                    pieza_path = os.path.join(folder_path, f)
                    
        faltantes = []
        if not resumen_path: faltantes.append("RESUMEN DE TRABAJO.pdf")
        if not nido_path: faltantes.append("DETALLE DEL NIDO DE CABEZAL.pdf")
        if not pieza_path: faltantes.append("DETALLE DE LA PIEZA.pdf")
        if faltantes:
            raise FileNotFoundError(f"No se encontraron en la carpeta: {', '.join(faltantes)}")
            
        doc_res = fitz.open(resumen_path)
        doc_nido = fitz.open(nido_path)
        doc_pieza = fitz.open(pieza_path)
    else:
        if not pdf_resumen_bytes or not pdf_nido_bytes or not pdf_pieza_bytes:
            raise ValueError("Se deben proporcionar los 3 archivos PDF.")
            
        doc_res = fitz.open(stream=pdf_resumen_bytes, filetype="pdf")
        doc_nido = fitz.open(stream=pdf_nido_bytes, filetype="pdf")
        doc_pieza = fitz.open(stream=pdf_pieza_bytes, filetype="pdf")

    # -------------------------------------------------------------
    # 1. PARSEAR DETALLE DE LA PIEZA (Catálogo de Nombres Completos)
    # -------------------------------------------------------------
    pieza_catalog = {}
    for page in doc_pieza:
        for block in page.get_text("blocks"):
            for line in block[4].split("\n"):
                line = line.strip()
                if "-(" in line:
                    short = line.split("-(")[0].strip()
                    pieza_catalog[short] = line

    # -------------------------------------------------------------
    # 2. PARSEAR RESUMEN DE TRABAJO (Nidos y Cortes / Hojas)
    # -------------------------------------------------------------
    nidos_hojas = []
    resumen_p1_text = doc_res[0].get_text()
    lines_p1 = [l.strip() for l in resumen_p1_text.split("\n") if l.strip()]

    of_title = "OF PRODUCIDA"
    for l in lines_p1[:6]:
        if l.startswith("OF ") or l.startswith("0F "):
            of_title = l
            break

    for page in doc_res:
        blocks = page.get_text("blocks")
        for b in blocks:
            b_text = b[4].strip()
            lines = [l.strip() for l in b_text.split("\n") if l.strip()]
            for i in range(len(lines)):
                if lines[i].isdigit() and i+1 < len(lines) and lines[i+1].isdigit():
                    if i+2 < len(lines) and "%" in lines[i+2]:
                        n_id = int(lines[i])
                        n_cortes = int(lines[i+1])
                        n_label = f"N{n_id:02d}"
                        if not any(item[0] == n_label for item in nidos_hojas):
                            nidos_hojas.append((n_label, n_cortes))

    nidos_hojas.sort(key=lambda x: int(x[0][1:]))

    # -------------------------------------------------------------
    # 3. PARSEAR DETALLE DEL NIDO DE CABEZAL (Piezas por Nido)
    # -------------------------------------------------------------
    piezas_rows = []
    material_cabezal = ""
    for b in doc_nido[0].get_text("blocks"):
        t = b[4].strip()
        if "Material:" in t:
            parts = t.split("Material:", 1)
            material_cabezal = parts[1].strip().split("\n")[0].strip()

    current_nido_num = 1
    for page_idx, page in enumerate(doc_nido):
        blocks = page.get_text("blocks")

        # Detectar si esta página inicia un nuevo nido o es continuación de páginas anteriores
        has_nido_header = False
        for b in blocks:
            if b[1] < 150 and "Nido:" in b[4]:
                has_nido_header = True
                break

        if has_nido_header:
            for b in blocks:
                if b[1] < 130 and "de" in b[4] and "ProNest" not in b[4] and "Detalle" not in b[4]:
                    m = re.search(r'(\d+)\s+de\s+(\d+)', b[4])
                    if m:
                        current_nido_num = int(m.group(1))
                        break
                m_direct = re.search(r'Nido:\s*(\d+)', b[4], re.IGNORECASE)
                if m_direct and b[1] < 150:
                    current_nido_num = int(m_direct.group(1))
                    break

        nido_label = f"N{current_nido_num:02d}"
        for b in blocks:
            text = b[4].strip()
            lines = [l.strip() for l in text.split("\n") if l.strip()]
            if len(lines) >= 4 and ("mm" in lines[0] or "kg" in lines[-1]):
                seq_str = lines[1]
                if "-" in seq_str:
                    p = seq_str.split("-")
                    try:
                        qty = int(p[1]) - int(p[0]) + 1
                    except:
                        continue
                elif seq_str.isdigit():
                    qty = 1
                else:
                    continue

                piece_raw = lines[3]
                piece_short = piece_raw.split("-(")[0].strip() if "-(" in piece_raw else piece_raw.split("(")[0].strip()
                piece_short = re.sub(r'\s*-\s*$', '', piece_short)
                full_name = pieza_catalog.get(piece_short, piece_raw)

                found = False
                for r in piezas_rows:
                    if r[0] == nido_label and r[1] == piece_short:
                        r[3] += qty
                        found = True
                        break
                if not found:
                    piezas_rows.append([nido_label, piece_short, full_name, qty])

    # -------------------------------------------------------------
    # 4. HOJA ORDEN: DETERMINACIÓN DE PARÁMETROS
    # -------------------------------------------------------------
    calibre_str = "Cal 12"
    m_cal = re.search(r'(\d+)\s*ga', material_cabezal, re.IGNORECASE)
    if not m_cal:
        m_cal = re.search(r'cal[.\s]*(\d+)', of_title, re.IGNORECASE)
    if m_cal:
        calibre_str = f"Cal {m_cal.group(1)}"

    fecha_prod = datetime.datetime.now()

    po_str = "N/A"
    m_po = re.search(r'(PO\s*\d+|OC\s*\d+[-\w]+|\d{4}-\d{4})', of_title, re.IGNORECASE)
    if m_po:
        po_str = m_po.group(1).strip()

    if params_orden is None:
        params_orden = {}

    nombre_proyecto = params_orden.get("nombre_proyecto") or of_title
    programador = params_orden.get("programador") or "BRYAN MANCINAS"
    orden_fab = params_orden.get("orden_fab") or of_title
    po = params_orden.get("po") or po_str
    desc_pronest = params_orden.get("desc_pronest") or of_title
    calibre = params_orden.get("calibre") or calibre_str
    prioridad = params_orden.get("prioridad") if params_orden.get("prioridad") is not None else 1
    cliente_proyecto = params_orden.get("cliente_proyecto") or "POR DEFINIR"

    orden_row = [
        nombre_proyecto,
        fecha_prod,
        programador,
        orden_fab,
        po,
        desc_pronest,
        calibre,
        prioridad,
        cliente_proyecto
    ]

    # -------------------------------------------------------------
    # 5. CONSTRUCCIÓN DEL WORKBOOK EXCEL (.xlsx)
    # -------------------------------------------------------------
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    font_header = Font(name="Calibri", size=11, bold=True, color="111111")
    fill_header = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")

    # Hoja 1: Orden
    ws_ord = wb.create_sheet(title="Orden")
    ws_ord.views.sheetView[0].showGridLines = True
    headers_ord = [
        'Nombre del Proyecto *', 'Fecha de Producción', 'Programador *',
        'Orden de Fabricación', 'PO', 'Descripción de OF Pronest',
        'Calibre', 'PRIORIDAD', 'Nombre del Proyecto de Cliente'
    ]
    ws_ord.append(headers_ord)
    ws_ord.append(orden_row)

    for col_idx in range(1, len(headers_ord) + 1):
        cell = ws_ord.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
    ws_ord.cell(row=2, column=2).number_format = 'yyyy-mm-dd'

    ws_ord.column_dimensions['A'].width = 28
    ws_ord.column_dimensions['B'].width = 20
    ws_ord.column_dimensions['C'].width = 22
    ws_ord.column_dimensions['D'].width = 48
    ws_ord.column_dimensions['E'].width = 20
    ws_ord.column_dimensions['F'].width = 48
    ws_ord.column_dimensions['G'].width = 14
    ws_ord.column_dimensions['H'].width = 14
    ws_ord.column_dimensions['I'].width = 28

    # Hoja 2: Nidos
    ws_nid = wb.create_sheet(title="Nidos")
    ws_nid.views.sheetView[0].showGridLines = True
    headers_nid = ['NIDO', 'HOJAS']
    ws_nid.append(headers_nid)
    for row in nidos_hojas:
        ws_nid.append(list(row))
    for col_idx in range(1, 3):
        cell = ws_nid.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
    ws_nid.column_dimensions['A'].width = 14
    ws_nid.column_dimensions['B'].width = 14

    # Hoja 3: Piezas
    ws_pie = wb.create_sheet(title="Piezas")
    ws_pie.views.sheetView[0].showGridLines = True
    headers_pie = ['NIDO', 'No. PIEZA', 'NOMBRE DE PIEZA', 'CANTIDAD']
    ws_pie.append(headers_pie)
    for row in piezas_rows:
        ws_pie.append(row)
    for col_idx in range(1, 5):
        cell = ws_pie.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
    ws_pie.column_dimensions['A'].width = 12
    ws_pie.column_dimensions['B'].width = 35
    ws_pie.column_dimensions['C'].width = 58
    ws_pie.column_dimensions['D'].width = 14

    excel_io = io.BytesIO()
    wb.save(excel_io)
    excel_io.seek(0)

    # Convert to DataFrames for display in Streamlit
    df_orden = pd.DataFrame([orden_row], columns=headers_ord)
    df_nidos = pd.DataFrame(nidos_hojas, columns=headers_nid)
    df_piezas = pd.DataFrame(piezas_rows, columns=headers_pie)

    stats = {
        "of_title": of_title,
        "folder_path": folder_path,
        "calibre": calibre,
        "total_nidos": len(nidos_hojas),
        "total_piezas_registros": len(piezas_rows),
        "total_piezas_fisicas": sum(r[3] for r in piezas_rows),
        "df_orden": df_orden,
        "df_nidos": df_nidos,
        "df_piezas": df_piezas,
        "excel_bytes": excel_io.getvalue()
    }
    return wb, stats

generar_excel_of_pronest = procesar_of_pronest
